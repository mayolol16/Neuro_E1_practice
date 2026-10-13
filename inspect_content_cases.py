import os
from pptx import Presentation

folder = "/Users/markyoussef_1/Downloads/NeuroE1ShEET"

content_pps = [
    "1-DevelcontentANS - Copy.pptx",
    "2-APSynAnesthContentANS - Copy (1).pptx",
    "3-VentHisoContentANS.pptx",
    "4-MGContentANS.pptx",
    "5-SpCrdcontentANS - Copy.pptx",
    "6-PaincontentANS - Copy.pptx",
    "7-TrigemcontentANS.pptx",
    "8-HeadachecontentANS.pptx",
    "9-TraumaContentANS.pptx",
    "Anesesthesia Content ANS.pptx",
    "MeningitisContentAnswers.pptx"
]

case_pps = [
    "1-DevelCASEANS - Copy.pptx",
    "4-NMJCASEANS.pptx",
    "5-SpCrdCASEANS - Copy.pptx",
    "6-PainCASEANS.pptx",
    "7-TrigemCASEANS.pptx",
    "9-TraumaCASEANS.pptx",
    "Anesesthesia CASE ANS.pptx",
    "MeningitisCaseAnswers..pptx"
]

print("="*60)
print("INSPECTING CONTENT POWERPOINTS")
print("="*60)
for cp in content_pps:
    path = os.path.join(folder, cp)
    prs = Presentation(path)
    print(f"\n>>> FILE: {cp} (Total slides: {len(prs.slides)})")
    for idx, slide in enumerate(prs.slides):
        texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        header = texts[0].replace('\n', ' ') if texts else "[No text]"
        full = " // ".join(texts)[:200].replace('\n', ' ')
        print(f"  Slide {idx+1}: {full}")

print("\n" + "="*60)
print("INSPECTING CASE POWERPOINTS")
print("="*60)
for cp in case_pps:
    path = os.path.join(folder, cp)
    prs = Presentation(path)
    print(f"\n>>> FILE: {cp} (Total slides: {len(prs.slides)})")
    for idx, slide in enumerate(prs.slides):
        texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        header = texts[0].replace('\n', ' ') if texts else "[No text]"
        full = " // ".join(texts)[:200].replace('\n', ' ')
        print(f"  Slide {idx+1}: {full}")
