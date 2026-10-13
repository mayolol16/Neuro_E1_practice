import os
from pptx import Presentation
import re

folder = "/Users/markyoussef_1/Downloads/NeuroE1ShEET"

content_pps = [
    "1-DevelcontentANS - Copy.pptx",
    "2-APSynAnesthContentANS - Copy (1).pptx",
    "3-VentHisoContentANS.pptx",
    "4-MGContentANS.pptx",
    "5-SpCrdcontentANS - Copy.pptx",
    "6-PaincontentANS - Copy.pptx",
    "7-TrigemcontentANS.pptx",
    "8-HeadachecontentANS.pptx",
    "9-TraumaContentANS.pptx",
    "Anesesthesia Content ANS.pptx",
    "MeningitisContentAnswers.pptx"
]

case_pps = [
    "1-DevelCASEANS - Copy.pptx",
    "4-NMJCASEANS.pptx",
    "5-SpCrdCASEANS - Copy.pptx",
    "6-PainCASEANS.pptx",
    "7-TrigemCASEANS.pptx",
    "9-TraumaCASEANS.pptx",
    "Anesesthesia CASE ANS.pptx",
    "MeningitisCaseAnswers..pptx"
]

def analyze_questions_in_file(fname):
    prs = Presentation(os.path.join(folder, fname))
    slide_texts = []
    for idx, s in enumerate(prs.slides):
        texts = [p.text.strip() for shape in s.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs if p.text.strip()]
        slide_texts.append((idx + 1, texts))
    return slide_texts

print("=== EXACT QUESTION INVENTORY ===")
for group_name, file_list in [("CONTENT POWERPOINTS", content_pps), ("CASE POWERPOINTS", case_pps)]:
    print(f"\n#################################################")
    print(f" {group_name}")
    print(f"#################################################")
    for f in file_list:
        slides = analyze_questions_in_file(f)
        print(f"\n--- {f} ---")
        for snum, txts in slides:
            if snum == 1 and ("content questions" in " ".join(txts).lower() or "case questions" in " ".join(txts).lower()):
                continue # Title slide
            joined = " \n ".join(txts)
            # Find lines starting with numbers
            print(f"  Slide {snum}:")
            for t in txts:
                if re.match(r'^\d+[\.\)]', t) or any(t.lower().startswith(x) for x in ['case', 'question', 'compare', 'describe', 'explain', 'what', 'list', 'how', 'which', 'define']):
                    print(f"    • {t[:120]}")
