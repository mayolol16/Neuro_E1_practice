import os
from pptx import Presentation
import fitz # PyMuPDF
import re

folder = "/Users/markyoussef_1/Downloads/NeuroE1ShEET"

def inspect_pptx(filepath):
    prs = Presentation(filepath)
    slides_data = []
    for idx, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        texts.append(t)
        slide_text = "\n".join(texts)
        slides_data.append((idx + 1, slide_text))
    return slides_data

def inspect_pdf(filepath):
    doc = fitz.open(filepath)
    pages_data = []
    for idx, page in enumerate(doc):
        t = page.get_text().strip()
        pages_data.append((idx + 1, t))
    return pages_data

print("="*60)
print("1. DEEP INSPECTION OF LECTURES")
print("="*60)

lectures = [
    "CNS Anatomy 2026.pdf",
    "CNS Development- student.pptx",
    "IMC 535 - Ventricles student.pptx",
    "Neuromuscular Junction Disorders.pptx",
    "Neurotransmission- student.pptx",
    "SMBS MS2 lecture 10-9-26.pptx",
    "Somatosensory pathway- ALS student.pptx",
    "Somatosensory pathway- DCML student.pptx",
    "Somatosensory pathway- Trigeminal student.pptx"
]

for lec in lectures:
    path = os.path.join(folder, lec)
    print(f"\n>>> LECTURE: {lec}")
    if lec.endswith('.pdf'):
        pages = inspect_pdf(path)
        print(f"Total pages: {len(pages)}")
        # Look for objectives
        for pnum, text in pages[:10]:
            if any(k in text.lower() for k in ["objective", "goal", "learning outcome", "competenc"]):
                print(f"  [Page {pnum} MATCH]:\n{text[:600]}\n---")
    else:
        slides = inspect_pptx(path)
        print(f"Total slides: {len(slides)}")
        for snum, text in slides[:10]:
            if any(k in text.lower() for k in ["objective", "goal", "learning outcome", "competenc"]):
                print(f"  [Slide {snum} MATCH]:\n{text[:600]}\n---")
