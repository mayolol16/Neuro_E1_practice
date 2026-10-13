import os
from pptx import Presentation
import re

folder = "/Users/markyoussef_1/Downloads/NeuroE1ShEET"

def get_slide_text(slide):
    texts = []
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    texts.append(txt)
    return "\n".join(texts)

files_to_check = [
    ("1-DevelcontentANS - Copy.pptx", "CONTENT"),
    ("2-APSynAnesthContentANS - Copy (1).pptx", "CONTENT"),
    ("3-VentHisoContentANS.pptx", "CONTENT"),
    ("4-MGContentANS.pptx", "CONTENT"),
    ("5-SpCrdcontentANS - Copy.pptx", "CONTENT"),
    ("6-PaincontentANS - Copy.pptx", "CONTENT"),
    ("7-TrigemcontentANS.pptx", "CONTENT"),
    ("8-HeadachecontentANS.pptx", "CONTENT"),
    ("9-TraumaContentANS.pptx", "CONTENT"),
    ("Anesesthesia Content ANS.pptx", "CONTENT"),
    ("MeningitisContentAnswers.pptx", "CONTENT"),
    ("1-DevelCASEANS - Copy.pptx", "CASE"),
    ("4-NMJCASEANS.pptx", "CASE"),
    ("5-SpCrdCASEANS - Copy.pptx", "CASE"),
    ("6-PainCASEANS.pptx", "CASE"),
    ("7-TrigemCASEANS.pptx", "CASE"),
    ("9-TraumaCASEANS.pptx", "CASE"),
    ("Anesesthesia CASE ANS.pptx", "CASE"),
    ("MeningitisCaseAnswers..pptx", "CASE")
]

for fname, ftype in files_to_check:
    prs = Presentation(os.path.join(folder, fname))
    print(f"\n{'='*70}")
    print(f"[{ftype}] {fname} ({len(prs.slides)} slides)")
    print(f"{'='*70}")
    for idx, slide in enumerate(prs.slides):
        txt = get_slide_text(slide)
        lines = [l.strip() for l in txt.split('\n') if l.strip()]
        header = lines[0] if lines else "[EMPTY]"
        # print first few lines of each slide
        print(f"  Slide {idx+1}:")
        for l in lines[:4]:
            print(f"    {l[:110]}")
